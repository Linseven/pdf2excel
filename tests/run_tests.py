#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PDF2Excel 端到端测试：生成样例 PDF → 转换 → 断言 Excel 内容。

运行：.venv/bin/python tests/run_tests.py
样例 PDF 保留在 tests/samples/ 下，可用于手动测试 GUI。
"""

from __future__ import annotations

import os
import shutil
import subprocess
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Table, TableStyle

from converter import ConversionError, convert, probe_pdf

# 注册 Adobe CID 中文字体，确保中文可被正确嵌入与提取
pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
CJK = "STSong-Light"
CJK_STYLE = ParagraphStyle("cjk", fontName=CJK, fontSize=12, leading=18)
CJK_TITLE = ParagraphStyle("cjk_title", fontName=CJK, fontSize=16, leading=24)

SAMPLES = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")
OUT = os.path.join(SAMPLES, "out")

PASS, FAIL = 0, 0


def check(name: str, cond: bool, detail: str = "") -> None:
    global PASS, FAIL
    if cond:
        PASS += 1
        print(f"  ✓ {name}")
    else:
        FAIL += 1
        print(f"  ✗ {name} {detail}")


# ------------------------------------------------------------------ 样例 PDF 生成

DATA = [
    ["编号", "名称", "数量", "单价"],
    ["A-1001", "机械键盘", "12", "¥329.00"],
    ["A-1002", "无线鼠标", "45", "¥99.50"],
    ["A-1003", "4K 显示器", "8", "¥2,199.00"],
    ["A-1004", "USB-C 扩展坞", "23", "¥459.00"],
]


def make_bordered(path: str, pages: int = 1) -> None:
    doc = SimpleDocTemplate(path, pagesize=A4)
    story = []
    for p in range(pages):
        if p:
            story.append(PageBreak())
        story.append(Paragraph(f"订单明细（第 {p + 1} 批）", CJK_TITLE))
        story.append(Table(DATA, colWidths=[70, 120, 50, 90], style=TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.8, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E0E7FF")),
            ("FONTNAME", (0, 0), (-1, -1), CJK),
            ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
        ])))
    doc.build(story)


def make_borderless(path: str) -> None:
    doc = SimpleDocTemplate(path, pagesize=A4)
    doc.build([Table(DATA, colWidths=[70, 120, 50, 90], style=TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), CJK),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
    ]))])


def make_textonly(path: str) -> None:
    doc = SimpleDocTemplate(path, pagesize=A4)
    doc.build([
        Paragraph("会议纪要", CJK_TITLE),
        Paragraph("时间：2026 年 10 月 14 日 上午 10:00", CJK_STYLE),
        Paragraph("地点：3 号会议室", CJK_STYLE),
        Paragraph("与会人：张三、李四、王五", CJK_STYLE),
        Paragraph("议题：关于第四季度产品发布节奏的讨论。经协商，各项交付节点保持不变，"
                  "市场部需在两周内完成物料准备。", CJK_STYLE),
    ])


def make_encrypted(path: str) -> None:
    from pypdf import PdfWriter

    plain = path + ".plain.pdf"
    make_bordered(plain)
    w = PdfWriter(clone_from=plain)
    w.encrypt("secret123")
    with open(path, "wb") as f:
        w.write(f)
    os.remove(plain)


def ensure_samples() -> None:
    os.makedirs(SAMPLES, exist_ok=True)
    shutil.rmtree(OUT, ignore_errors=True)
    os.makedirs(OUT, exist_ok=True)
    make_bordered(os.path.join(SAMPLES, "bordered.pdf"))
    make_borderless(os.path.join(SAMPLES, "borderless.pdf"))
    make_textonly(os.path.join(SAMPLES, "textonly.pdf"))
    make_bordered(os.path.join(SAMPLES, "multipage.pdf"), pages=5)
    make_encrypted(os.path.join(SAMPLES, "encrypted.pdf"))


# ------------------------------------------------------------------ 测试用例


def sheet_rows(ws):
    return [[c.value for c in row] for row in ws.iter_rows() if any(c.value not in (None, "") for c in row)]


def test_bordered_auto():
    print("[1] 线框表格 · 自动模式")
    res = convert(os.path.join(SAMPLES, "bordered.pdf"), xlsx_path=os.path.join(OUT, "t1.xlsx"))
    ws = load_workbook(res["output"]).active
    rows = sheet_rows(ws)
    check("提取到 5 行 4 列", len(rows) == 5 and all(len(r) >= 4 for r in rows), str(rows))
    check("表头正确", rows[0][:4] == ["编号", "名称", "数量", "单价"], str(rows[0]))
    check("数据正确", rows[2][:4] == ["A-1002", "无线鼠标", "45", "¥99.50"], str(rows[2]))
    check("表头加粗", ws.cell(1, 1).font.bold is True)
    check("冻结首行", ws.freeze_panes == "A2")
    check("列宽已设置", ws.column_dimensions["B"].width is not None)


def test_borderless_text_mode():
    print("[2] 无框表格 · text 模式")
    res = convert(os.path.join(SAMPLES, "borderless.pdf"), xlsx_path=os.path.join(OUT, "t2.xlsx"), mode="text")
    ws = load_workbook(res["output"]).active
    rows = sheet_rows(ws)
    check("提取到表格", len(rows) == 5 and max(len(r) for r in rows) >= 4, str(rows))
    check("内容正确", any(r[:2] == ["A-1002", "无线鼠标"] for r in rows), str(rows))


def test_auto_fallback_to_text():
    print("[3] 无框表格 · auto 自动回退")
    res = convert(os.path.join(SAMPLES, "borderless.pdf"), xlsx_path=os.path.join(OUT, "t3.xlsx"), mode="auto")
    rows = sheet_rows(load_workbook(res["output"]).active)
    check("自动回退提取成功", len(rows) == 5, str(rows))


def test_page_range_and_sheets():
    print("[4] 多页 PDF · 页码范围 + 工作表模式")
    res = convert(os.path.join(SAMPLES, "multipage.pdf"), xlsx_path=os.path.join(OUT, "t4.xlsx"),
                  pages="2-4", sheet_mode="pages")
    wb = load_workbook(res["output"])
    check("生成 3 个工作表", len(wb.sheetnames) == 3, str(wb.sheetnames))
    check("只处理了 3 页", res["pages"] == 3)
    rows = sheet_rows(wb[wb.sheetnames[0]])
    check("每页表格 5 行（表头+4 数据；页标题不属于表格，不提取）", len(rows) == 5, str(len(rows)))

    res2 = convert(os.path.join(SAMPLES, "multipage.pdf"), xlsx_path=os.path.join(OUT, "t4b.xlsx"),
                   sheet_mode="merge")
    wb2 = load_workbook(res2["output"])
    check("合并模式仅 1 个工作表", len(wb2.sheetnames) == 1, str(wb2.sheetnames))
    rows2 = sheet_rows(wb2.active)
    check("5 页共 25 行", len(rows2) == 25, str(len(rows2)))


def test_text_fallback():
    print("[5] 纯文本 PDF · 文本回退")
    res = convert(os.path.join(SAMPLES, "textonly.pdf"), xlsx_path=os.path.join(OUT, "t5.xlsx"))
    rows = sheet_rows(load_workbook(res["output"]).active)
    check("提取到文本行", len(rows) >= 4, str(rows))
    check("包含关键内容", any("会议纪要" in str(r[0]) for r in rows), str(rows[:2]))
    check("产生回退提示", any("纯文本" in w for w in res["warnings"]))


def test_encrypted():
    print("[6] 加密 PDF")
    info = probe_pdf(os.path.join(SAMPLES, "encrypted.pdf"))
    check("探测到加密", info["encrypted"] is True)

    try:
        convert(os.path.join(SAMPLES, "encrypted.pdf"), xlsx_path=os.path.join(OUT, "t6.xlsx"))
        check("无密码时报错", False)
    except ConversionError as e:
        check("无密码时报错", e.code == "password", str(e))

    try:
        convert(os.path.join(SAMPLES, "encrypted.pdf"), xlsx_path=os.path.join(OUT, "t6.xlsx"), password="wrong")
        check("错误密码时报错", False)
    except ConversionError as e:
        check("错误密码时报错", e.code == "password", str(e))

    res = convert(os.path.join(SAMPLES, "encrypted.pdf"), xlsx_path=os.path.join(OUT, "t6.xlsx"),
                  password="secret123")
    rows = sheet_rows(load_workbook(res["output"]).active)
    check("正确密码转换成功", rows[0][:4] == ["编号", "名称", "数量", "单价"], str(rows[0]))


def test_unique_output():
    print("[7] 输出文件名去重")
    p1 = convert(os.path.join(SAMPLES, "bordered.pdf"), out_dir=OUT)["output"]
    p2 = convert(os.path.join(SAMPLES, "bordered.pdf"), out_dir=OUT)["output"]
    check("自动追加序号", p1 != p2 and "(1)" in p2, p2)


def test_cli():
    print("[8] CLI 模式")
    out = os.path.join(OUT, "cli.xlsx")
    r = subprocess.run(
        [sys.executable, os.path.join(ROOT, "main.py"), "--cli",
         os.path.join(SAMPLES, "bordered.pdf"), "-o", out],
        capture_output=True, text=True, timeout=120)
    check("退出码为 0", r.returncode == 0, r.stderr[-300:] if r.returncode else "")
    check("生成文件", os.path.exists(out))


def main() -> int:
    print("生成样例 PDF …")
    ensure_samples()
    print()
    test_bordered_auto()
    test_borderless_text_mode()
    test_auto_fallback_to_text()
    test_page_range_and_sheets()
    test_text_fallback()
    test_encrypted()
    test_unique_output()
    test_cli()
    print(f"\n结果：{PASS} 通过 / {FAIL} 失败")
    return 1 if FAIL else 0


if __name__ == "__main__":
    sys.exit(main())
