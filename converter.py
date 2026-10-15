# -*- coding: utf-8 -*-
"""PDF → Excel 核心转换引擎。

纯逻辑模块，不依赖任何 GUI 框架，可被 CLI / GUI / 测试共同复用。
"""

from __future__ import annotations

import os
import re
import unicodedata

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# pdfminer.six 的密码异常（不同版本类名不同，做兼容导入）
try:
    from pdfminer.pdfdocument import PDFPasswordIncorrect as _PasswordError
except ImportError:  # pragma: no cover
    try:
        from pdfminer.pdfdocument import PDFPasswordError as _PasswordError
    except ImportError:
        _PasswordError = ()  # except 子句可安全使用的占位


def _is_password_error(exc: Exception) -> bool:
    """判断异常是否为“PDF 需要密码 / 密码错误”。

    pdfplumber 会把 pdfminer 的原始异常包装为 PdfminerException（原始异常挂在
    __cause__ / __context__ 上且消息为空），因此需要沿异常链检查类型名。
    """
    cur, depth = exc, 0
    while cur is not None and depth < 6:
        if isinstance(cur, _PasswordError):
            return True
        if "password" in type(cur).__name__.lower() or "password" in str(cur).lower():
            return True
        cur = getattr(cur, "__cause__", None) or getattr(cur, "__context__", None)
        depth += 1
    return False

__all__ = [
    "ConversionError",
    "ConversionCancelled",
    "convert",
    "probe_pdf",
    "unique_output_path",
    "parse_page_range",
]

# ------------------------------------------------------------------ 异常


class ConversionError(Exception):
    """携带面向用户的友好信息的转换失败。code: password / unreadable / empty / generic"""

    def __init__(self, message: str, code: str = "generic"):
        super().__init__(message)
        self.code = code


class ConversionCancelled(Exception):
    """用户主动取消。"""


# ------------------------------------------------------------------ 表格提取参数

# 线框表格：依赖 PDF 中的矢量线条划分单元格
_LINES_SETTINGS = {
    "vertical_strategy": "lines",
    "horizontal_strategy": "lines",
}

# 无框表格：依据文字的列对齐/间距推断单元格边界
_TEXT_SETTINGS = {
    "vertical_strategy": "text",
    "horizontal_strategy": "text",
    "text_x_tolerance": 2,
    "text_y_tolerance": 3,
}

# ------------------------------------------------------------------ Excel 样式

_HEADER_FILL = PatternFill("solid", fgColor="EEF2FF")
_HEADER_FONT = Font(bold=True, color="1F2937")
_THIN = Side(style="thin", color="D9DEE7")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_ALIGN = Alignment(vertical="center")

_COL_MIN_WIDTH = 8
_COL_MAX_WIDTH = 60


# ------------------------------------------------------------------ 工具函数


def _friendly(exc: Exception) -> str:
    msg = str(exc).strip().splitlines()
    return msg[0] if msg else exc.__class__.__name__


def _clean_cell(value) -> str:
    """单元格清洗：None → ''，合并内部换行，去首尾空白。"""
    if value is None:
        return ""
    s = str(value).replace("\u00a0", " ")
    s = re.sub(r"\s*\n\s*", " ", s)
    return s.strip()


def _disp_width(text: str) -> int:
    """近似显示宽度：全角字符按 2 计。"""
    return sum(2 if unicodedata.east_asian_width(ch) in ("F", "W") else 1 for ch in text)


def _sanitize_sheet_title(name: str, fallback: str = "Sheet1") -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", name).strip()
    return (name or fallback)[:31] or fallback


def parse_page_range(spec: str | None, total_pages: int) -> list[int]:
    """解析 '1-5,8,11-13' 形式的页码范围（1-based），返回升序去重页码。"""
    spec = (spec or "").strip()
    if not spec:
        return list(range(1, total_pages + 1))
    nums: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a > b:
                a, b = b, a
            nums.update(range(a, b + 1))
        elif part.isdigit():
            nums.add(int(part))
        else:
            raise ConversionError(f"页码范围格式无效：“{part}”（示例：1-5,8,11-13）")
    pages = sorted(p for p in nums if 1 <= p <= total_pages)
    if not pages:
        raise ConversionError(f"页码范围超出文档页数（该 PDF 共 {total_pages} 页）")
    return pages


def unique_output_path(out_dir: str, stem: str) -> str:
    """在 out_dir 下为 stem 生成不冲突的 .xlsx 路径：name.xlsx / name (1).xlsx ..."""
    os.makedirs(out_dir, exist_ok=True)
    stem = stem or "output"
    candidate = os.path.join(out_dir, stem + ".xlsx")
    n = 1
    while os.path.exists(candidate):
        candidate = os.path.join(out_dir, f"{stem} ({n}).xlsx")
        n += 1
    return candidate


def _valid_text_tables(tables: list) -> list:
    """text 策略的防幻觉过滤：至少 2 行、且存在 ≥2 列的行。"""
    out = []
    for t in tables:
        if len(t) >= 2 and max(len(r) for r in t) >= 2:
            rows = [r for r in t if any(_clean_cell(c) for c in r)]
            if len(rows) >= 2:
                out.append(rows)
    return out


def _extract_tables(page, mode: str) -> list:
    """按模式提取当前页的表格（已清洗：去除全空行）。"""
    if mode == "lines":
        raw = page.extract_tables(_LINES_SETTINGS)
    elif mode == "text":
        return _valid_text_tables(page.extract_tables(_TEXT_SETTINGS))
    else:  # auto：先按线框提取，失败再按文字对齐提取
        raw = page.extract_tables(_LINES_SETTINGS)
        if not any(any(r) for r in raw):
            return _valid_text_tables(page.extract_tables(_TEXT_SETTINGS))
    out = []
    for t in raw:
        rows = [r for r in t if any(_clean_cell(c) for c in r)]
        if rows:
            out.append(rows)
    return out


def _write_table(ws, rows: list, start_row: int, widths: dict) -> int:
    """把一个表格写入工作表，返回下一个可用行号。widths: {列号: 最大显示宽}。"""
    n_rows = len(rows)
    n_cols = max(len(r) for r in rows)
    for i, row in enumerate(rows):
        for j in range(n_cols):
            v = _clean_cell(row[j]) if j < len(row) else ""
            cell = ws.cell(row=start_row + i, column=j + 1, value=v)
            cell.border = _BORDER
            cell.alignment = _ALIGN
            if i == 0:  # 首行视为表头
                cell.font = _HEADER_FONT
                cell.fill = _HEADER_FILL
            widths[j + 1] = max(widths.get(j + 1, 0), _disp_width(v))
    return start_row + n_rows


def _apply_widths(ws, widths: dict) -> None:
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(_COL_MIN_WIDTH, min(w + 2, _COL_MAX_WIDTH))


def _write_lines(ws, lines: list, start_row: int, widths: dict) -> int:
    """纯文本回退：每行文本写入一个单元格。"""
    for i, line in enumerate(lines):
        cell = ws.cell(row=start_row + i, column=1, value=line)
        cell.alignment = _ALIGN
        widths[1] = max(widths.get(1, 0), _disp_width(line))
    return start_row + len(lines)


# ------------------------------------------------------------------ 对外接口


def probe_pdf(path: str, password: str | None = None) -> dict:
    """探测 PDF 基本信息（供文件列表展示）：页数 / 是否加密 / 是否可读。"""
    path = os.path.abspath(path)
    info = {
        "path": path,
        "name": os.path.basename(path),
        "size": os.path.getsize(path) if os.path.exists(path) else 0,
        "pages": None,
        "encrypted": False,
        "error": None,
    }
    if not os.path.exists(path):
        info["error"] = "文件不存在"
        return info
    try:
        with pdfplumber.open(path, password=password or "") as pdf:
            info["pages"] = len(pdf.pages)
    except Exception as e:  # noqa: BLE001 —— 统一转为友好信息
        if _is_password_error(e):
            info["encrypted"] = True
        else:
            info["error"] = f"无法读取：{_friendly(e)}"
    return info


def convert(
    pdf_path: str,
    xlsx_path: str | None = None,
    out_dir: str | None = None,
    mode: str = "auto",
    pages: str | None = "",
    sheet_mode: str = "merge",
    password: str | None = None,
    progress_cb=None,
    cancel_event=None,
) -> dict:
    """将单个 PDF 转换为 Excel。

    参数
    ----
    pdf_path    : 输入 PDF
    xlsx_path   : 输出 xlsx 路径；为 None 时按 out_dir/同名.xlsx 自动生成
    mode        : auto（自动识别）| lines（线框表格）| text（无框表格）
    pages       : 页码范围，如 "1-5,8"；空 = 全部
    sheet_mode  : merge（所有页合并到一个工作表）| pages（每页一个工作表）
    password    : 加密 PDF 的打开密码
    progress_cb : 回调 fn(done, total)，每处理完一页触发一次
    cancel_event: threading.Event，置位后中止转换（抛 ConversionCancelled）

    返回 dict: {output, pages, tables, rows, warnings}
    """
    pdf_path = os.path.abspath(pdf_path)
    if not os.path.isfile(pdf_path):
        raise ConversionError(f"文件不存在：{pdf_path}")
    if xlsx_path and os.path.isdir(xlsx_path):
        out_dir, xlsx_path = xlsx_path, None
    if not xlsx_path:
        stem = os.path.splitext(os.path.basename(pdf_path))[0]
        xlsx_path = unique_output_path(out_dir or os.path.dirname(pdf_path), stem)

    try:
        pdf = pdfplumber.open(pdf_path, password=password or "")
    except Exception as e:  # noqa: BLE001
        if _is_password_error(e):
            if password:
                raise ConversionError("密码错误，请重新输入", code="password") from e
            raise ConversionError("该 PDF 已加密，请输入打开密码", code="password") from e
        raise ConversionError(f"无法读取 PDF 文件：{_friendly(e)}", code="unreadable") from e

    warnings: list[str] = []
    tables_total = 0
    rows_total = 0

    with pdf:
        total = len(pdf.pages)
        page_nums = parse_page_range(pages, total)

        wb = Workbook()
        stem = os.path.splitext(os.path.basename(pdf_path))[0]
        merge_ws = wb.active
        merge_ws.title = _sanitize_sheet_title(stem)
        merge_widths: dict = {}
        next_row = 1

        for idx, pno in enumerate(page_nums):
            if cancel_event is not None and cancel_event.is_set():
                raise ConversionCancelled()

            page = pdf.pages[pno - 1]
            page_tables = _extract_tables(page, mode)

            if page_tables:
                # ---- 表格页：写入一个（merge）或当前页自己的（pages）工作表
                if sheet_mode == "pages":
                    ws = wb.create_sheet(_sanitize_sheet_title(f"第 {pno} 页", f"Page {pno}"))
                    widths: dict = {}
                    row = 1
                    for table in page_tables:
                        row = _write_table(ws, table, row, widths) + 1  # 表格间空一行
                        tables_total += 1
                        rows_total += len(table)
                    ws.freeze_panes = "A2"
                    _apply_widths(ws, widths)
                else:
                    for table in page_tables:
                        next_row = _write_table(merge_ws, table, next_row, merge_widths) + 1
                        tables_total += 1
                        rows_total += len(table)
                    if merge_ws.freeze_panes is None and next_row > 1:
                        merge_ws.freeze_panes = "A2"
            elif len(page.chars) == 0:
                # ---- 既无表格也无任何文字对象：疑似扫描件 / 纯图片页
                warnings.append(f"第 {pno} 页未检测到文字内容（可能为扫描件或图片页，需 OCR 才能提取）")
                note = f"〔第 {pno} 页无可提取内容（疑似扫描件）〕"
                if sheet_mode == "pages":
                    ws = wb.create_sheet(_sanitize_sheet_title(f"第 {pno} 页", f"Page {pno}"))
                    _write_lines(ws, [note], 1, {})
                else:
                    next_row = _write_lines(merge_ws, [note], next_row, merge_widths) + 1
            else:
                # ---- 有文字但没识别出表格：文本回退
                text = page.extract_text() or ""
                lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
                if lines:
                    warnings.append(f"第 {pno} 页未检测到表格，已按纯文本提取")
                    if sheet_mode == "pages":
                        ws = wb.create_sheet(_sanitize_sheet_title(f"第 {pno} 页", f"Page {pno}"))
                        _write_lines(ws, lines, 1, {})
                    else:
                        next_row = _write_lines(merge_ws, lines, next_row, merge_widths) + 1
                    rows_total += len(lines)
                else:
                    warnings.append(f"第 {pno} 页内容为空，已跳过")

            if progress_cb is not None:
                progress_cb(idx + 1, len(page_nums))

        if tables_total == 0 and rows_total == 0:
            raise ConversionError(
                "未能从该 PDF 提取到任何表格或文字（文件可能是扫描件，需要 OCR 支持）",
                code="empty",
            )

        if sheet_mode != "pages":
            _apply_widths(merge_ws, merge_widths)
        else:
            merge_ws.title = "_empty_"
            wb.remove(merge_ws)  # 移除未使用的默认工作表

        wb.save(xlsx_path)

    return {
        "output": os.path.abspath(xlsx_path),
        "pages": len(page_nums),
        "tables": tables_total,
        "rows": rows_total,
        "warnings": warnings,
    }
